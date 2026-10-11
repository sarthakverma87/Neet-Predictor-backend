"""Owner-only helper: build the analytics Excel file (sheets: Leads, Events, Funnel).

    python export_excel.py                    # everything -> neet_analytics.xlsx
    python export_excel.py --real             # leave out test leads and test events
    python export_excel.py --out my_file.xlsx

The sheets copy the layout of the reference project's Google Sheet, with these differences:
  * Leads: name and phone are masked (Asha Rao -> A****o, 9876543210 -> 9****0), and the utm_source,
    utm_medium, utm_campaign and duplicate_phone columns are left out.
  * Funnel: only page_view, form_start, form_submit and result_view (no cta_click, no per-UTM table).
  * Events: unchanged from the reference.
The database itself keeps the real values; only this file is masked.
"""
import sqlite3
import sys
from contextlib import closing
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

import database

DEFAULT_OUT = Path(__file__).parent / "neet_analytics.xlsx"

LEAD_HEADERS = ["timestamp", "name", "phone", "score", "category", "state", "predicted_rank", "is_test",
                "score_band", "device", "safe_count", "moderate_count", "reach_count"]
EVENT_HEADERS = ["timestamp", "event", "utm_source", "utm_medium", "utm_campaign", "device", "score_band",
                 "category", "state", "safe_count", "moderate_count", "reach_count", "cta_type"]
FUNNEL_EVENTS = ["page_view", "form_start", "form_submit", "result_view"]


def mask_name(name):
    """First letter + **** + last letter. Anything shorter than 2 characters is fully hidden."""
    name = (name or "").strip()
    return name[0] + "****" + name[-1] if len(name) >= 2 else "****"


def mask_phone(phone):
    """First digit + **** + last digit."""
    phone = (phone or "").strip()
    return phone[0] + "****" + phone[-1] if len(phone) >= 2 else "****"


def _bool(flag):
    return "TRUE" if flag else "FALSE"


def _write_sheet(ws, headers, rows, widths):
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in rows:
        ws.append(["" if v is None else v for v in row])
    # Values come from visitors (names, UTM tags). openpyxl turns any text starting with "=" into a formula,
    # so force every text cell to stay plain text.
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            if isinstance(cell.value, str):
                cell.data_type = "s"
    ws.freeze_panes = "A2"
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w


def build_workbook(conn, real_only=False):
    test_filter = " WHERE is_test = 0" if real_only else ""

    leads = conn.execute(
        "SELECT created_at, name, mobile, score, category, home_state, estimated_air, is_test, score_band, "
        "device, safe_count, moderate_count, reach_count FROM leads" + test_filter + " ORDER BY id").fetchall()
    lead_rows = [(r[0], mask_name(r[1]), mask_phone(r[2]), r[3], r[4], r[5], r[6], _bool(r[7]),
                  r[8], r[9], r[10], r[11], r[12]) for r in leads]

    events = conn.execute(
        "SELECT created_at, event, utm_source, utm_medium, utm_campaign, device, score_band, category, state, "
        "safe_count, moderate_count, reach_count, cta_type FROM events" + test_filter + " ORDER BY id").fetchall()

    counts = {name: 0 for name in FUNNEL_EVENTS}
    for row in events:
        if row[1] in counts:
            counts[row[1]] += 1
    funnel_rows = [(name, counts[name]) for name in FUNNEL_EVENTS]

    wb = Workbook()
    ws = wb.active
    ws.title = "Leads"
    _write_sheet(ws, LEAD_HEADERS, lead_rows, [22, 12, 10, 8, 10, 24, 14, 8, 12, 10, 11, 14, 12])
    _write_sheet(wb.create_sheet("Events"), EVENT_HEADERS, events, [22, 14, 14, 14, 16, 10, 12, 10, 24, 11, 14, 12, 18])
    _write_sheet(wb.create_sheet("Funnel"), ["event", "count"], funnel_rows, [16, 10])
    return wb, len(lead_rows), len(events)


def main(argv):
    real_only = "--real" in argv
    out = Path(argv[argv.index("--out") + 1]) if "--out" in argv else DEFAULT_OUT
    database.init_db()  # makes the tables if the server has never run (and upgrades an older database)
    with closing(sqlite3.connect(database.db_path())) as conn:
        wb, n_leads, n_events = build_workbook(conn, real_only)
    try:
        wb.save(out)
    except PermissionError:
        print(f"Could not write {out}. Close it in Excel and run this again.")
        return 1
    print(f"Saved {out}  ({n_leads} lead(s), {n_events} event(s)){'  [test data excluded]' if real_only else ''}")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
