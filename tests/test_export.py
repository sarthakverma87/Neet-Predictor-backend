"""Tests for the analytics Excel export. Needs only openpyxl.   Run:  python -m unittest tests.test_export -v"""
import os
import sqlite3
import sys
import tempfile
import unittest
from contextlib import closing
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
_tmp = tempfile.TemporaryDirectory(ignore_cleanup_errors=True)
os.environ["DB_PATH"] = os.path.join(_tmp.name, "export.db")

import database  # noqa: E402
import export_excel  # noqa: E402

database.init_db()
database.save_lead("Asha Rao", "9876543210", 600, "OBC", "Delhi", 1300, "google", "cpc", "neet26", False,
                   device="mobile", score_band="600-649", safe_count=5, moderate_count=7, reach_count=9)
database.save_lead("Test Person", "9123456780", 500, "UR", "Bihar", 9000, None, None, None, True)
database.save_lead("=SUM(1)+2", "9000000001", 450, "SC", "Assam", 20000)
for ev in ("page_view", "page_view", "page_view", "form_start", "form_start", "form_submit", "result_view"):
    database.save_event(ev, utm_source="google", device="mobile")
database.save_event("share_click", cta_type="whatsapp", device="mobile")
database.save_event("cta_click", cta_type="counselling_call", is_test=True)
database.save_event("page_view", is_test=True)


def _workbook(real_only=False):
    with closing(sqlite3.connect(database.db_path())) as conn:
        wb, _, _ = export_excel.build_workbook(conn, real_only)
    return wb


def _rows(ws):
    return [[c.value for c in row] for row in ws.iter_rows()]


class MaskTests(unittest.TestCase):
    def test_name_and_phone_masks(self):
        self.assertEqual(export_excel.mask_name("Asha Rao"), "A****o")
        self.assertEqual(export_excel.mask_name("  Li  "), "L****i")
        self.assertEqual(export_excel.mask_name("A"), "****")
        self.assertEqual(export_excel.mask_phone("9876543210"), "9****0")
        self.assertEqual(export_excel.mask_phone(""), "****")


class WorkbookTests(unittest.TestCase):
    def test_three_sheets_in_order(self):
        self.assertEqual(_workbook().sheetnames, ["Leads", "Events", "Funnel"])

    def test_leads_sheet_columns_and_masking(self):
        rows = _rows(_workbook()["Leads"])
        self.assertEqual(rows[0], ["timestamp", "name", "phone", "score", "category", "state", "predicted_rank",
                                   "is_test", "score_band", "device", "safe_count", "moderate_count", "reach_count"])
        for banned in ("utm_source", "utm_medium", "utm_campaign", "duplicate_phone"):
            self.assertNotIn(banned, rows[0])
        asha = rows[1]
        self.assertEqual((asha[1], asha[2], asha[3], asha[4], asha[5], asha[6]), ("A****o", "9****0", 600, "OBC", "Delhi", 1300))
        self.assertEqual((asha[7], asha[8], asha[9], asha[10], asha[11], asha[12]), ("FALSE", "600-649", "mobile", 5, 7, 9))
        self.assertEqual(rows[2][7], "TRUE")  # the test lead
        flat = " ".join(str(v) for row in rows for v in row)
        for secret in ("Asha", "Rao", "9876543210", "Person", "google", "cpc"):
            self.assertNotIn(secret, flat)

    def test_funnel_has_only_the_four_steps(self):
        rows = _rows(_workbook()["Funnel"])
        self.assertEqual(rows, [["event", "count"], ["page_view", 4], ["form_start", 2],
                                ["form_submit", 1], ["result_view", 1]])

    def test_events_sheet_keeps_every_event_and_all_columns(self):
        rows = _rows(_workbook()["Events"])
        self.assertEqual(rows[0], ["timestamp", "event", "utm_source", "utm_medium", "utm_campaign", "device",
                                   "score_band", "category", "state", "safe_count", "moderate_count",
                                   "reach_count", "cta_type"])
        self.assertEqual(len(rows) - 1, 10)
        self.assertIn("share_click", [r[1] for r in rows[1:]])  # not in the funnel, but still an event
        self.assertEqual(rows[1][2], "google")                   # UTM columns are kept on Events

    def test_real_only_drops_test_rows(self):
        wb = _workbook(real_only=True)
        self.assertEqual(len(_rows(wb["Leads"])) - 1, 2)
        self.assertEqual(len(_rows(wb["Events"])) - 1, 8)
        self.assertEqual(_rows(wb["Funnel"])[1], ["page_view", 3])

    def test_text_starting_with_equals_stays_text_not_a_formula(self):
        ws = _workbook()["Leads"]
        cell = ws.cell(row=4, column=2)  # name "=SUM(1)+2" -> "=****2"
        self.assertEqual(cell.value, "=****2")
        self.assertEqual(cell.data_type, "s")

    def test_file_saves_and_reloads(self):
        out = Path(_tmp.name) / "out.xlsx"
        self.assertEqual(export_excel.main(["--out", str(out)]), 0)
        from openpyxl import load_workbook
        self.assertEqual(load_workbook(out).sheetnames, ["Leads", "Events", "Funnel"])


class MigrationTests(unittest.TestCase):
    def test_old_database_is_upgraded_in_place(self):
        old = os.path.join(_tmp.name, "old.db")
        with closing(sqlite3.connect(old)) as conn, conn:
            conn.execute("""CREATE TABLE leads (id INTEGER PRIMARY KEY AUTOINCREMENT, name TEXT NOT NULL,
                mobile TEXT NOT NULL, score INTEGER NOT NULL, category TEXT NOT NULL, home_state TEXT NOT NULL,
                estimated_air INTEGER, utm_source TEXT, utm_medium TEXT, utm_campaign TEXT,
                is_test INTEGER NOT NULL DEFAULT 0, submit_count INTEGER NOT NULL DEFAULT 1,
                created_at TEXT NOT NULL, updated_at TEXT NOT NULL)""")
            conn.execute("INSERT INTO leads (name, mobile, score, category, home_state, created_at, updated_at) "
                         "VALUES ('Old Lead','9000000009',400,'UR','Goa','2026-01-01','2026-01-01')")
        previous = os.environ["DB_PATH"]
        os.environ["DB_PATH"] = old
        try:
            database.init_db()
            with closing(sqlite3.connect(old)) as conn:
                cols = {r[1] for r in conn.execute("PRAGMA table_info(leads)")}
                wb, n, _ = export_excel.build_workbook(conn)
        finally:
            os.environ["DB_PATH"] = previous
        self.assertTrue({"score_band", "device", "safe_count", "moderate_count", "reach_count"} <= cols)
        self.assertEqual(n, 1)
        self.assertEqual(_rows(wb["Leads"])[1][1], "O****d")


if __name__ == "__main__":
    unittest.main()
